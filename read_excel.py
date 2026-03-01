import openpyxl

wb = openpyxl.load_workbook('xperp_menu_list.xlsx', data_only=True)
ws = wb.sheetnames[0]
ws = wb[ws]

print('=== F 컬럼(설명) 없는 항목 ===')
missing = []
for r in range(2, ws.max_row + 1):
    a = ws.cell(r, 1).value  # 대분류
    b = ws.cell(r, 2).value  # 중분류
    c = ws.cell(r, 3).value  # 소분류
    f = ws.cell(r, 6).value  # 정의/기능
    if any([a, b, c]) and not f:
        missing.append((r, a, b, c))
        print(f'  {r}행: {a} > {b} > {c}')

print(f'\n총 {len(missing)}개 항목 설명 없음')

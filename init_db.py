"""
XPERP 매뉴얼 챗봇 - Neo4j 초기화 스크립트
- xperp_menu_list.xlsx에서 메뉴 데이터 로드
- Menu / SubMenu / MenuItem 노드 구조 생성
- Gemini 임베딩 생성 및 menu_vector_index 생성

실행 방법:
    uv pip install openpyxl
    .venv\\Scripts\\python.exe init_db.py
"""

import os
import openpyxl
from dotenv import load_dotenv
import neo4j
from gemini_embedder import GeminiEmbedder

load_dotenv()

NEO4J_URI = os.getenv("NEO4J_URI", "neo4j://localhost:7687")
NEO4J_PASSWORD = os.getenv("NEO4J_PASSWORD", "password")
GOOGLE_API_KEY = os.getenv("GOOGLE_API_KEY")
INDEX_NAME = "menu_vector_index"
EMBEDDING_DIM = 3072  # gemini-embedding-001 차원 수


def load_menu_data(filepath="xperp_menu_list.xlsx"):
    """엑셀에서 메뉴 데이터 로드 (설명 있는 항목만)"""
    wb = openpyxl.load_workbook(filepath, data_only=True)
    ws = wb[wb.sheetnames[0]]

    items = []
    skipped = 0
    for r in range(2, ws.max_row + 1):
        a = ws.cell(r, 1).value  # 대분류
        b = ws.cell(r, 2).value  # 중분류
        c = ws.cell(r, 3).value  # 소분류
        f = ws.cell(r, 6).value  # 정의/기능

        if not a and not b:
            continue
        if not f:
            skipped += 1
            continue

        items.append({
            "main_menu": str(a).strip() if a else None,
            "sub_menu": str(b).strip() if b else None,
            "menu_item": str(c).strip() if c else None,
            "description": str(f).strip(),
        })

    print(f"  로드 완료: {len(items)}개 항목 (설명 없어 스킵: {skipped}개)")
    return items


def clear_database(session):
    """기존 데이터 전체 삭제"""
    session.run("MATCH (n) DETACH DELETE n")


def create_constraints(session):
    """유니크 제약 조건 생성"""
    constraints = [
        "CREATE CONSTRAINT menu_name IF NOT EXISTS FOR (m:Menu) REQUIRE m.name IS UNIQUE",
        "CREATE CONSTRAINT submenu_id IF NOT EXISTS FOR (s:SubMenu) REQUIRE s.id IS UNIQUE",
        "CREATE CONSTRAINT menuitem_id IF NOT EXISTS FOR (i:MenuItem) REQUIRE i.id IS UNIQUE",
    ]
    for constraint in constraints:
        try:
            session.run(constraint)
        except Exception as e:
            print(f"  제약 조건 이미 존재: {e}")


def insert_data(session, embedder, items):
    """메뉴 데이터 삽입 및 임베딩 생성"""
    total = len(items)
    for i, item in enumerate(items, 1):
        main = item["main_menu"]
        sub = item["sub_menu"]
        menu_item = item["menu_item"]
        desc = item["description"]

        label = f"{main} > {sub} > {menu_item}" if menu_item else f"{main} > {sub}"
        print(f"[{i}/{total}] {label}", end=" ... ", flush=True)

        # Menu 노드 생성
        session.run("MERGE (m:Menu {name: $name})", name=main)

        # 임베딩 생성
        embedding = embedder.embed_query(desc)

        if menu_item:
            # 3-level: Menu → SubMenu → MenuItem
            sub_id = f"{main}::{sub}"
            item_id = f"{main}::{sub}::{menu_item}"
            menu_path = f"{main} > {sub} > {menu_item}"

            session.run(
                "MERGE (s:SubMenu {id: $id}) SET s.name = $name, s.main_menu = $main",
                id=sub_id, name=sub, main=main
            )
            session.run(
                """
                MATCH (m:Menu {name: $main})
                MATCH (s:SubMenu {id: $sub_id})
                MERGE (m)-[:HAS_SUBMENU]->(s)
                """,
                main=main, sub_id=sub_id
            )
            session.run(
                """
                MERGE (i:MenuItem {id: $id})
                SET i.name = $name,
                    i.description = $description,
                    i.main_menu = $main,
                    i.sub_menu = $sub,
                    i.menu_path = $menu_path,
                    i.embedding = $embedding
                """,
                id=item_id, name=menu_item, description=desc,
                main=main, sub=sub, menu_path=menu_path, embedding=embedding
            )
            session.run(
                """
                MATCH (s:SubMenu {id: $sub_id})
                MATCH (i:MenuItem {id: $item_id})
                MERGE (s)-[:HAS_ITEM]->(i)
                """,
                sub_id=sub_id, item_id=item_id
            )
        else:
            # 2-level: Menu → MenuItem (직접 연결)
            item_id = f"{main}::{sub}"
            menu_path = f"{main} > {sub}"

            session.run(
                """
                MERGE (i:MenuItem {id: $id})
                SET i.name = $name,
                    i.description = $description,
                    i.main_menu = $main,
                    i.sub_menu = $sub,
                    i.menu_path = $menu_path,
                    i.embedding = $embedding
                """,
                id=item_id, name=sub, description=desc,
                main=main, sub=sub, menu_path=menu_path, embedding=embedding
            )
            session.run(
                """
                MATCH (m:Menu {name: $main})
                MATCH (i:MenuItem {id: $item_id})
                MERGE (m)-[:HAS_ITEM]->(i)
                """,
                main=main, item_id=item_id
            )

        print(f"✅ (dim={len(embedding)})")

    print(f"\n총 {total}개 항목 삽입 완료")


def create_vector_index(session):
    """MenuItem 벡터 인덱스 생성"""
    try:
        session.run(f"DROP INDEX {INDEX_NAME} IF EXISTS")
    except Exception:
        pass

    session.run(
        f"""
        CREATE VECTOR INDEX {INDEX_NAME} IF NOT EXISTS
        FOR (i:MenuItem) ON i.embedding
        OPTIONS {{
            indexConfig: {{
                `vector.dimensions`: {EMBEDDING_DIM},
                `vector.similarity_function`: 'cosine'
            }}
        }}
        """
    )


def verify_data(session):
    """데이터 검증"""
    menus = session.run("MATCH (m:Menu) RETURN count(m) AS c").single()["c"]
    submenus = session.run("MATCH (s:SubMenu) RETURN count(s) AS c").single()["c"]
    items = session.run("MATCH (i:MenuItem) RETURN count(i) AS c").single()["c"]
    embedded = session.run(
        "MATCH (i:MenuItem) WHERE i.embedding IS NOT NULL RETURN count(i) AS c"
    ).single()["c"]
    indexes = session.run("SHOW INDEXES WHERE name = $name", name=INDEX_NAME).data()

    print(f"  Menu 노드     : {menus}개")
    print(f"  SubMenu 노드  : {submenus}개")
    print(f"  MenuItem 노드 : {items}개 (임베딩: {embedded}개)")
    print(f"  벡터 인덱스   : {'✅ 존재' if indexes else '❌ 없음'}")


def main():
    print("=" * 60)
    print("  XPERP 매뉴얼 챗봇 - Neo4j 초기화")
    print("=" * 60)
    print(f"  Neo4j URI : {NEO4J_URI}")
    print(f"  Index     : {INDEX_NAME}")
    print("=" * 60)

    print("\n📋 엑셀 데이터 로드 중...")
    items = load_menu_data()

    print("\n📡 Gemini 임베더 초기화 중...")
    embedder = GeminiEmbedder(model="gemini-embedding-001", api_key=GOOGLE_API_KEY)
    print("✅ 임베더 초기화 완료")

    print(f"\n🔌 Neo4j 연결 중: {NEO4J_URI}")
    driver = neo4j.GraphDatabase.driver(NEO4J_URI, auth=("neo4j", NEO4J_PASSWORD))
    try:
        driver.verify_connectivity()
        print("✅ Neo4j 연결 성공\n")
    except Exception as e:
        print(f"❌ Neo4j 연결 실패: {e}")
        return

    with driver.session() as session:
        print("🗑️  기존 데이터 삭제 중...")
        clear_database(session)
        print("✅ 삭제 완료\n")

        print("📌 제약 조건 생성 중...")
        create_constraints(session)
        print("✅ 제약 조건 완료\n")

        print("📥 데이터 삽입 및 임베딩 생성 중...")
        insert_data(session, embedder, items)

        print("\n🔍 벡터 인덱스 생성 중...")
        create_vector_index(session)
        print("✅ 인덱스 생성 완료\n")

        print("📊 데이터 검증:")
        verify_data(session)

    driver.close()
    print("\n🎉 초기화 완료! 이제 서버를 실행하세요:")
    print("   .venv\\Scripts\\uvicorn app:app --host 0.0.0.0 --port 8800 --reload")


if __name__ == "__main__":
    main()
